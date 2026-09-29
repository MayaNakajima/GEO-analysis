# -*- coding: utf-8 -*-
"""
施策管理表（Google参考値_施策管理.xlsx）を作る／新しい設問の行を追加する
────────────────────────────────────────────────────────────────
Box の GEO-analysis フォルダに置き、関係者が「担当者・状態・実施内容」を記入する Excel。
generate.py（gsc_reader.load_actions）が読み、Google参考値タブに反映する。

    python tools/make_action_sheet.py              # 無ければ作る（ある場合は何もしない）
    python tools/make_action_sheet.py --add-missing # 既存の表に、まだ行のない設問を末尾に追加する

依存：openpyxl（Anaconda に同梱）。generate.py 本体は標準ライブラリのみのまま。
設問は GEO の monitoring/config/questions.json・questions_set2.json、担当の初期値は config.json の domain_owner。
"""

import argparse
import json
import os
import sys

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, HERE)
import generate  # noqa: E402  （config の読み込みとパスの決め方を共通にする）

HEAD = ["設問ID", "領域（参照）", "設問（参照）", "担当（事業）", "担当者", "状態", "期限", "実施日", "実施内容", "対象ページURL", "メモ"]
WIDTH = [17, 18, 48, 11, 12, 10, 12, 12, 44, 36, 24]
NOTE = ("記入方法：該当する設問ID の行に「担当者」「状態」「期限」「実施日」「実施内容」「対象ページURL」を書いて保存。"
        "同じ設問で次の施策をしたら行を追加し、同じ設問ID を書く（下の行ほど新しい記録）。"
        "「担当（事業）」を書き換えると、その設問の担当が変わる。編集は1人ずつ、保存したら閉じる。")


def questions(cfg):
    mon = os.path.dirname(os.path.dirname(os.path.abspath(cfg["results_dir"])))   # monitoring/
    out = []
    for f in ("questions.json", "questions_set2.json"):
        p = os.path.join(mon, "config", f)
        if os.path.exists(p):
            out += json.load(open(p, encoding="utf-8"))
    return out


def main():
    ap = argparse.ArgumentParser(description="施策管理表を作る")
    ap.add_argument("--config", default=generate.DEFAULT_CONFIG)
    ap.add_argument("--out", default=None, help="出力先（省略時は config の action_xlsx／share_dirs の最初）")
    ap.add_argument("--add-missing", action="store_true", help="既存の表に、行のない設問を追加する")
    a = ap.parse_args()

    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    cfg = generate.load_config(a.config)
    path = a.out or generate.action_xlsx_path(cfg)
    if not path:
        sys.exit("[ERROR] 出力先が決まりません（config.json の share_dirs か action_xlsx を設定）")
    owner = cfg.get("domain_owner") or {}
    qs = questions(cfg)

    if os.path.exists(path) and not a.add_missing:
        print(f"[SKIP] すでにあります（上書きしません）: {path}")
        print("       新しい設問の行を足すときは --add-missing を付けて実行してください。")
        return
    if a.add_missing:
        if not os.path.exists(path):
            sys.exit(f"[ERROR] 見つかりません: {path}")
        wb = load_workbook(path)
        ws = wb["施策管理"] if "施策管理" in wb.sheetnames else wb.active
        have = {str(ws.cell(r, 1).value or "").strip() for r in range(3, ws.max_row + 1)}
        n = 0
        for q in qs:
            if q["id"] not in have:
                ws.append([q["id"], q.get("domain_label", ""), q.get("question", ""),
                           owner.get(q.get("domain_label", ""), "未割当")])
                n += 1
        wb.save(path)
        print(f"[OK] {n} 行を追加しました: {path}")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "施策管理"
    ws.append([NOTE])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEAD))
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(size=9, color="5F6B7A")
    ws.row_dimensions[1].height = 42
    ws.append(HEAD)
    for i, w in enumerate(WIDTH, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
        c = ws.cell(2, i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1D4F91" if i <= 3 else "2E7D5B")
        c.alignment = Alignment(vertical="center")
    for q in qs:
        ws.append([q["id"], q.get("domain_label", ""), q.get("question", ""),
                   owner.get(q.get("domain_label", ""), "未割当")])
    last = ws.max_row + 200            # 行を追加して記入できるよう、入力規則は少し先まで
    dv = DataValidation(type="list", formula1='"未着手,対応中,完了,見送り"', allow_blank=True,
                        showErrorMessage=True, errorTitle="状態", error="未着手／対応中／完了／見送り から選んでください")
    ws.add_data_validation(dv)
    dv.add(f"F3:F{last}")
    for r in range(3, ws.max_row + 1):       # 空の行にセルを作らない（行の追加位置がずれるため）
        for col in ("G", "H"):
            ws[f"{col}{r}"].number_format = "yyyy-mm-dd"
        ws[f"C{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        for col in "AB":
            ws[f"{col}{r}"].font = Font(color="5F6B7A")
    ws.freeze_panes = "D3"
    ws.auto_filter.ref = f"A2:{chr(64 + len(HEAD))}{ws.max_row}"
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    print(f"[OK] 作成しました（{len(qs)} 問）: {path}")


if __name__ == "__main__":
    main()
